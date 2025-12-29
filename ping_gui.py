import sqlite3
import subprocess
import platform
import threading
import time
import re
import queue
import concurrent.futures
from datetime import datetime

from PySide6 import QtCore, QtWidgets
from PySide6.QtCharts import QChart, QChartView, QLineSeries, QScatterSeries, QValueAxis
from PySide6.QtGui import QPainter

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

DEFAULT_PING_INTERVAL = 600
DEFAULT_MONITOR_INTERVAL = 1
DEFAULT_PING_TIMEOUT = 1
DEFAULT_BATCH_SIZE = 50
DEFAULT_MAX_WORKERS = 20
STATS_REFRESH_INTERVAL = 5


def clean_address(address):
    address = address.strip()
    if not address:
        raise ValueError("Адрес не может быть пустым.")

    ipv4_pattern = r"^(?:(?:25[0-5]|2[0-4][0-9]|[01]?[0-9][0-9]?)\.){3}(?:25[0-5]|2[0-4][0-9]|[01]?[0-9][0-9]?)$"
    ipv6_pattern = r"^(?:[0-9a-fA-F]{1,4}:){7}[0-9a-fA-F]{1,4}$"
    domain_pattern = r"^[a-zA-Z0-9](?:[a-zA-Z0-9-]{0,61}[a-zA-Z0-9])?(?:\.[a-zA-Z0-9](?:[a-zA-Z0-9-]{0,61}[a-zA-Z0-9])?)*$"

    if not (
        re.match(ipv4_pattern, address)
        or re.match(ipv6_pattern, address)
        or re.match(domain_pattern, address)
    ):
        raise ValueError(
            f"Некорректный адрес: {address}. Должен быть IPv4, IPv6 или доменным именем."
        )

    return address


def clean_group_name(group_name):
    group_name = group_name.strip()
    if not group_name:
        raise ValueError("Имя группы не может быть пустым.")
    if not re.match(r"^[a-zA-Z0-9_\sа-яА-ЯәңғұүқөһіӘҢҒҰҮҚӨҺІ\-\.\(\)]+$", group_name):
        raise ValueError("Некорректное имя группы. Используйте буквы, цифры и пробелы.")
    return group_name


class DatabaseManager:
    def __init__(self, db_name="monitoring_gui.db"):
        self.conn = sqlite3.connect(db_name, check_same_thread=False)
        self.cursor = self.conn.cursor()
        self.lock = threading.RLock()
        self.init_db()
        self.migrate_legacy_tables()

    def init_db(self):
        with self.lock:
            self.cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS groups (
                    group_name TEXT PRIMARY KEY,
                    active INTEGER DEFAULT 0
                )
            """
            )
            self.cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS hosts (
                    group_name TEXT NOT NULL,
                    address TEXT NOT NULL,
                    description TEXT,
                    subgroup TEXT,
                    ping_interval INTEGER DEFAULT 600,
                    last_ping_time TEXT,
                    offline_since TEXT,
                    PRIMARY KEY (group_name, address)
                )
            """
            )
            self.cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS ping_results (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    timestamp TEXT,
                    group_name TEXT,
                    subgroup TEXT,
                    address TEXT,
                    status TEXT,
                    latency REAL
                )
            """
            )
            self.cursor.execute(
                """
                CREATE INDEX IF NOT EXISTS idx_hosts_group
                ON hosts (group_name)
            """
            )
            self.cursor.execute(
                """
                CREATE INDEX IF NOT EXISTS idx_ping_results_group_address
                ON ping_results (group_name, address)
            """
            )
            self.cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS settings (
                    key TEXT PRIMARY KEY,
                    value TEXT
                )
            """
            )
            self.cursor.executemany(
                """
                INSERT OR IGNORE INTO settings (key, value)
                VALUES (?, ?)
            """,
                [
                    ("default_interval_seconds", str(DEFAULT_PING_INTERVAL)),
                    ("monitor_interval_seconds", str(DEFAULT_MONITOR_INTERVAL)),
                    ("ping_timeout_seconds", str(DEFAULT_PING_TIMEOUT)),
                    ("ping_batch_size", str(DEFAULT_BATCH_SIZE)),
                    ("max_ping_workers", str(DEFAULT_MAX_WORKERS)),
                ],
            )
            self.conn.commit()
            self.ensure_groups_schema()

    def ensure_groups_schema(self):
        self.cursor.execute("PRAGMA table_info(groups)")
        columns = [row[1] for row in self.cursor.fetchall()]
        if "active" not in columns:
            self.cursor.execute("ALTER TABLE groups ADD COLUMN active INTEGER DEFAULT 0")
            self.conn.commit()

    def migrate_legacy_tables(self):
        with self.lock:
            self.cursor.execute(
                """
                SELECT name FROM sqlite_master
                WHERE type='table' AND name LIKE 'hosts_%'
            """
            )
            legacy_tables = [row[0] for row in self.cursor.fetchall()]
            for table_name in legacy_tables:
                group_name = table_name[len("hosts_") :]
                if not group_name:
                    continue
                self.cursor.execute(
                    "INSERT OR IGNORE INTO groups (group_name, active) VALUES (?, 0)",
                    (group_name,),
                )
                self.cursor.execute(
                    f"SELECT address, description, subgroup, ping_interval, last_ping_time, offline_since FROM {table_name}"
                )
                rows = self.cursor.fetchall()
                if rows:
                    self.cursor.executemany(
                        """
                        INSERT OR REPLACE INTO hosts
                        (group_name, address, description, subgroup, ping_interval, last_ping_time, offline_since)
                        VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                        [
                            (
                                group_name,
                                row[0],
                                row[1],
                                row[2],
                                row[3],
                                row[4],
                                row[5],
                            )
                            for row in rows
                        ],
                    )
                self.cursor.execute(f"DROP TABLE IF EXISTS {table_name}")
            self.conn.commit()

    def create_group(self, group_name):
        group_name = clean_group_name(group_name)
        with self.lock:
            self.cursor.execute(
                "INSERT OR IGNORE INTO groups (group_name, active) VALUES (?, 0)",
                (group_name,),
            )
            self.conn.commit()
        return group_name

    def delete_group(self, group_name):
        group_name = clean_group_name(group_name)
        with self.lock:
            self.cursor.execute("DELETE FROM groups WHERE group_name = ?", (group_name,))
            self.cursor.execute("DELETE FROM hosts WHERE group_name = ?", (group_name,))
            self.cursor.execute(
                "DELETE FROM ping_results WHERE group_name = ?", (group_name,)
            )
            self.conn.commit()

    def get_groups(self):
        with self.lock:
            self.cursor.execute("SELECT group_name FROM groups")
            return [row[0] for row in self.cursor.fetchall()]

    def get_groups_with_status(self):
        with self.lock:
            self.cursor.execute("SELECT group_name, active FROM groups")
            return {row[0]: bool(row[1]) for row in self.cursor.fetchall()}

    def set_group_active(self, group_name, active):
        group_name = clean_group_name(group_name)
        with self.lock:
            self.cursor.execute(
                "UPDATE groups SET active = ? WHERE group_name = ?",
                (1 if active else 0, group_name),
            )
            self.conn.commit()

    def get_active_groups(self):
        with self.lock:
            self.cursor.execute("SELECT group_name FROM groups WHERE active = 1")
            return [row[0] for row in self.cursor.fetchall()]

    def get_setting(self, key, default=None):
        with self.lock:
            self.cursor.execute("SELECT value FROM settings WHERE key = ?", (key,))
            row = self.cursor.fetchone()
            return row[0] if row else default

    def set_setting(self, key, value):
        with self.lock:
            self.cursor.execute(
                """
                INSERT INTO settings (key, value)
                VALUES (?, ?)
                ON CONFLICT(key) DO UPDATE SET value = excluded.value
            """,
                (key, str(value)),
            )
            self.conn.commit()

    def add_host(self, group_name, address, description, subgroup=None, ping_interval=DEFAULT_PING_INTERVAL):
        group_name = clean_group_name(group_name)
        address = clean_address(address)
        if subgroup:
            subgroup = clean_group_name(subgroup)
        with self.lock:
            self.cursor.execute(
                """
                INSERT OR REPLACE INTO hosts
                (group_name, address, description, subgroup, ping_interval, last_ping_time, offline_since)
                VALUES (?, ?, ?, ?, ?, NULL, NULL)
            """,
                (group_name, address, description, subgroup, ping_interval),
            )
            self.conn.commit()

    def update_host(self, group_name, address, description, subgroup=None, ping_interval=DEFAULT_PING_INTERVAL):
        group_name = clean_group_name(group_name)
        address = clean_address(address)
        if subgroup:
            subgroup = clean_group_name(subgroup)
        with self.lock:
            self.cursor.execute(
                """
                UPDATE hosts
                SET description = ?, subgroup = ?, ping_interval = ?
                WHERE group_name = ? AND address = ?
            """,
                (description, subgroup, ping_interval, group_name, address),
            )
            self.conn.commit()

    def remove_host(self, group_name, address):
        group_name = clean_group_name(group_name)
        address = clean_address(address)
        with self.lock:
            self.cursor.execute(
                "DELETE FROM hosts WHERE group_name = ? AND address = ?",
                (group_name, address),
            )
            self.conn.commit()

    def get_hosts(self, group_name):
        group_name = clean_group_name(group_name)
        with self.lock:
            self.cursor.execute(
                """
                SELECT address, description, subgroup, ping_interval, last_ping_time, offline_since
                FROM hosts
                WHERE group_name = ?
            """,
                (group_name,),
            )
            return self.cursor.fetchall()

    def get_last_status(self, group_name, address):
        group_name = clean_group_name(group_name)
        address = clean_address(address)
        with self.lock:
            self.cursor.execute(
                """
                SELECT status FROM ping_results
                WHERE group_name = ? AND address = ?
                ORDER BY id DESC LIMIT 1
            """,
                (group_name, address),
            )
            row = self.cursor.fetchone()
            return row[0] if row else "Unknown"

    def get_recent_results(self, group_name, address, limit=50):
        group_name = clean_group_name(group_name)
        address = clean_address(address)
        with self.lock:
            self.cursor.execute(
                """
                SELECT timestamp, status, latency
                FROM ping_results
                WHERE group_name = ? AND address = ?
                ORDER BY id DESC
                LIMIT ?
            """,
                (group_name, address, limit),
            )
            return list(reversed(self.cursor.fetchall()))

    def get_last_status_change(self, group_name, address, limit=200):
        group_name = clean_group_name(group_name)
        address = clean_address(address)
        with self.lock:
            self.cursor.execute(
                """
                SELECT timestamp, status
                FROM ping_results
                WHERE group_name = ? AND address = ?
                ORDER BY id DESC
                LIMIT ?
            """,
                (group_name, address, limit),
            )
            rows = self.cursor.fetchall()
        if not rows:
            return None, None
        last_status = rows[0][1]
        last_change_time = rows[0][0]
        for timestamp, status in rows[1:]:
            if status != last_status:
                break
            last_change_time = timestamp
        return last_status, last_change_time

    def update_ping_time(self, group_name, address, status, commit=True):
        group_name = clean_group_name(group_name)
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with self.lock:
            if status == "Online":
                self.cursor.execute(
                    """
                    UPDATE hosts
                    SET last_ping_time = ?, offline_since = NULL
                    WHERE group_name = ? AND address = ?
                """,
                    (now, group_name, address),
                )
            else:
                self.cursor.execute(
                    """
                    SELECT offline_since FROM hosts
                    WHERE group_name = ? AND address = ?
                """,
                    (group_name, address),
                )
                row = self.cursor.fetchone()
                if row and row[0] is None:
                    self.cursor.execute(
                        """
                        UPDATE hosts
                        SET last_ping_time = ?, offline_since = ?
                        WHERE group_name = ? AND address = ?
                    """,
                        (now, now, group_name, address),
                    )
                else:
                    self.cursor.execute(
                        """
                        UPDATE hosts
                        SET last_ping_time = ?
                        WHERE group_name = ? AND address = ?
                    """,
                        (now, group_name, address),
                    )
            if commit:
                self.conn.commit()

    def log_results_batch(self, entries):
        if not entries:
            return
        with self.lock:
            self.cursor.executemany(
                """
                INSERT INTO ping_results
                (timestamp, group_name, subgroup, address, status, latency)
                VALUES (?, ?, ?, ?, ?, ?)
            """,
                [
                    (
                        entry["timestamp"],
                        entry["group_name"],
                        entry["subgroup"],
                        entry["address"],
                        entry["status"],
                        entry["latency"],
                    )
                    for entry in entries
                ],
            )
            for entry in entries:
                self.update_ping_time(
                    entry["group_name"], entry["address"], entry["status"], commit=False
                )
            self.conn.commit()

    def get_all_hosts_stats(self):
        with self.lock:
            self.cursor.execute(
                """
                SELECT h.group_name, h.address, h.offline_since, pr.status
                FROM hosts h
                LEFT JOIN (
                    SELECT pr.group_name, pr.address, pr.status
                    FROM ping_results pr
                    INNER JOIN (
                        SELECT group_name, address, MAX(id) AS max_id
                        FROM ping_results
                        GROUP BY group_name, address
                    ) latest
                    ON pr.id = latest.max_id
                ) pr
                ON pr.group_name = h.group_name AND pr.address = h.address
            """
            )
            rows = self.cursor.fetchall()
            self.cursor.execute("SELECT group_name FROM groups")
            all_groups = [row[0] for row in self.cursor.fetchall()]

        stats_by_group = {}
        for group, address, offline_since, status in rows:
            group_stats = stats_by_group.setdefault(
                group,
                {
                    "group": group,
                    "total": 0,
                    "online": 0,
                    "offline": 0,
                    "unknown": 0,
                    "status": "unknown",
                },
            )
            group_stats["total"] += 1

            if status == "Online":
                group_stats["online"] += 1
            elif status == "Offline":
                group_stats["offline"] += 1
            else:
                group_stats["unknown"] += 1

            if status == "Offline":
                if offline_since:
                    try:
                        offline_time = datetime.strptime(
                            offline_since, "%Y-%m-%d %H:%M:%S"
                        )
                        if (datetime.now() - offline_time).total_seconds() >= 3600:
                            group_stats["status"] = "critical"
                        elif group_stats["status"] != "critical":
                            group_stats["status"] = "warning"
                    except ValueError:
                        group_stats["status"] = "critical"
                elif group_stats["status"] != "critical":
                    group_stats["status"] = "warning"

        for group_stats in stats_by_group.values():
            if group_stats["status"] == "unknown":
                if group_stats["offline"] > 0:
                    group_stats["status"] = "warning"
                elif group_stats["total"] == 0:
                    group_stats["status"] = "healthy"
                else:
                    group_stats["status"] = "healthy"
        for group in all_groups:
            stats_by_group.setdefault(
                group,
                {
                    "group": group,
                    "total": 0,
                    "online": 0,
                    "offline": 0,
                    "unknown": 0,
                    "status": "healthy",
                },
            )
        return list(stats_by_group.values())

    def get_overall_status(self):
        stats = self.get_all_hosts_stats()
        critical_count = sum(1 for s in stats if s["status"] == "critical")
        warning_count = sum(1 for s in stats if s["status"] == "warning")

        if critical_count > 0:
            return "critical"
        if warning_count > 0:
            return "warning"
        return "healthy"

    def close(self):
        with self.lock:
            self.conn.close()


def ping_host_subprocess(address, timeout=1.0):
    param = "-n" if platform.system().lower() == "windows" else "-c"
    timeout_param = "-w" if platform.system().lower() == "windows" else "-W"

    timeout_val = str(int(timeout * 1000)) if platform.system().lower() == "windows" else str(
        int(timeout)
    )
    if platform.system().lower() != "windows":
        timeout_val = str(max(1, int(timeout)))

    command = ["ping", param, "1", timeout_param, timeout_val, address]

    try:
        startupinfo = None
        if platform.system().lower() == "windows":
            startupinfo = subprocess.STARTUPINFO()
            startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW

        start_time = time.time()
        result = subprocess.run(
            command,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            startupinfo=startupinfo,
            text=True,
        )
        end_time = time.time()

        duration = end_time - start_time

        if result.returncode == 0:
            return True, duration
        return False, None
    except Exception:
        return False, None


class PingApp(QtWidgets.QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Ping Monitor")
        self.resize(1200, 700)

        self.db = DatabaseManager()
        self.monitoring = False
        self.monitor_thread = None
        self.monitor_interval = DEFAULT_MONITOR_INTERVAL
        self.ping_timeout = DEFAULT_PING_TIMEOUT
        self.batch_size = DEFAULT_BATCH_SIZE
        self.max_workers = DEFAULT_MAX_WORKERS
        self.default_interval_seconds = DEFAULT_PING_INTERVAL
        self.current_group = None
        self.host_status_cache = {}
        self.ui_queue = queue.Queue()
        self.last_stats_refresh = 0
        self.selected_host = None
        self.settings_dialog = None
        self.add_host_dialog = None
        self.edit_host_dialog = None
        self.create_group_dialog_ref = None
        self.history_dialog = None

        self.load_settings()
        self.build_ui()
        self.load_groups()
        self.update_overall_status()

        self.ui_timer = QtCore.QTimer(self)
        self.ui_timer.timeout.connect(self.process_ui_queue)
        self.ui_timer.start(200)

        if not HAS_OPENPYXL:
            QtWidgets.QMessageBox.warning(
                self,
                "Внимание",
                "openpyxl не установлен. Импорт/экспорт Excel будет недоступен.",
            )

    def load_settings(self):
        self.default_interval_seconds = int(
            self.db.get_setting("default_interval_seconds", DEFAULT_PING_INTERVAL)
        )
        self.monitor_interval = float(
            self.db.get_setting("monitor_interval_seconds", DEFAULT_MONITOR_INTERVAL)
        )
        self.ping_timeout = float(
            self.db.get_setting("ping_timeout_seconds", DEFAULT_PING_TIMEOUT)
        )
        self.batch_size = int(self.db.get_setting("ping_batch_size", DEFAULT_BATCH_SIZE))
        self.max_workers = int(self.db.get_setting("max_ping_workers", DEFAULT_MAX_WORKERS))

    def build_ui(self):
        central = QtWidgets.QWidget()
        self.setCentralWidget(central)
        layout = QtWidgets.QVBoxLayout(central)
        layout.setContentsMargins(8, 8, 8, 8)
        layout.setSpacing(8)

        self.tabs = QtWidgets.QTabWidget()
        layout.addWidget(self.tabs)

        self.monitor_tab = QtWidgets.QWidget()
        self.stats_tab = QtWidgets.QWidget()
        self.tabs.addTab(self.monitor_tab, "Мониторинг")
        self.tabs.addTab(self.stats_tab, "Статистика")

        self.build_monitor_tab()
        self.build_stats_tab()

        status_layout = QtWidgets.QHBoxLayout()
        layout.addLayout(status_layout)
        status_layout.addWidget(QtWidgets.QLabel("Общий статус:"))
        self.overall_status_label = QtWidgets.QLabel("●")
        status_layout.addWidget(self.overall_status_label)
        status_layout.addStretch()
        self.status_label = QtWidgets.QLabel("Готов")
        status_layout.addWidget(self.status_label)

    def build_monitor_tab(self):
        layout = QtWidgets.QVBoxLayout(self.monitor_tab)
        layout.setSpacing(6)

        toolbar = QtWidgets.QHBoxLayout()
        layout.addLayout(toolbar)

        toolbar.addWidget(QtWidgets.QLabel("Группа:"))
        self.group_combo = QtWidgets.QComboBox()
        self.group_combo.currentTextChanged.connect(self.on_group_select)
        toolbar.addWidget(self.group_combo)

        toolbar.addWidget(self.create_button("Создать", self.show_create_group_dialog))
        toolbar.addWidget(self.create_button("Удалить", self.delete_group))
        toolbar.addSpacing(10)

        toolbar.addWidget(self.create_button("Добавить хост", self.show_add_host_dialog))
        toolbar.addWidget(self.create_button("Редактировать хост", self.show_edit_host_dialog))
        toolbar.addWidget(self.create_button("Удалить хост", self.remove_host))
        toolbar.addSpacing(10)

        toolbar.addWidget(self.create_button("Экспорт (Excel)", self.export_data))
        toolbar.addWidget(self.create_button("Импорт (Excel)", self.import_data))
        toolbar.addWidget(self.create_button("Настройки", self.show_settings_dialog))
        toolbar.addSpacing(10)

        toolbar.addWidget(QtWidgets.QLabel("Поиск:"))
        self.search_input = QtWidgets.QLineEdit()
        self.search_input.textChanged.connect(self.refresh_table)
        toolbar.addWidget(self.search_input)
        toolbar.addWidget(self.create_button("Очистить", self.clear_search))
        toolbar.addSpacing(10)

        toolbar.addWidget(QtWidgets.QLabel("Фильтр:"))
        self.filter_combo = QtWidgets.QComboBox()
        self.filter_combo.addItems([
            "all",
            "online",
            "offline_lt_1h",
            "offline_ge_1h",
            "unknown",
        ])
        self.filter_combo.currentTextChanged.connect(self.refresh_table)
        toolbar.addWidget(self.filter_combo)
        toolbar.addStretch()

        self.btn_start = QtWidgets.QPushButton("Старт")
        self.btn_start.clicked.connect(self.toggle_monitoring)
        toolbar.addWidget(self.btn_start)

        content = QtWidgets.QSplitter()
        content.setStretchFactor(0, 3)
        content.setStretchFactor(1, 1)
        layout.addWidget(content)

        self.host_table = QtWidgets.QTableWidget(0, 7)
        self.host_table.setHorizontalHeaderLabels(
            [
                "",
                "Адрес",
                "Описание",
                "Подгруппа",
                "Интервал (сек)",
                "Задержка (мс)",
                "Последняя проверка",
            ]
        )
        self.host_table.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
        self.host_table.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.host_table.verticalHeader().setVisible(False)
        self.host_table.horizontalHeader().setStretchLastSection(True)
        self.host_table.itemSelectionChanged.connect(self.on_host_select)
        content.addWidget(self.host_table)

        detail_widget = QtWidgets.QWidget()
        detail_layout = QtWidgets.QVBoxLayout(detail_widget)
        detail_layout.setContentsMargins(8, 8, 8, 8)
        detail_layout.setSpacing(8)
        content.addWidget(detail_widget)

        detail_layout.addWidget(self.section_label("Детали хоста"))
        self.detail_labels = {}
        for label in [
            "Адрес",
            "Описание",
            "Подгруппа",
            "Интервал",
            "Статус",
            "Задержка",
            "Последняя проверка",
            "Последняя смена",
        ]:
            row = QtWidgets.QHBoxLayout()
            row.addWidget(QtWidgets.QLabel(f"{label}:"))
            value = QtWidgets.QLabel("—")
            value.setTextInteractionFlags(QtCore.Qt.TextSelectableByMouse)
            row.addWidget(value)
            row.addStretch()
            detail_layout.addLayout(row)
            self.detail_labels[label] = value

        metrics_box = QtWidgets.QGroupBox("Метрики")
        metrics_layout = QtWidgets.QGridLayout(metrics_box)
        metrics_layout.setContentsMargins(8, 8, 8, 8)
        metrics_layout.setHorizontalSpacing(16)
        metrics_layout.setVerticalSpacing(8)
        metric_keys = [
            ("Мин", "min"),
            ("Макс", "max"),
            ("Средняя", "avg"),
            ("Джиттер", "jitter"),
            ("Потери", "loss"),
        ]
        self.metric_labels = {}
        for idx, (title, key) in enumerate(metric_keys):
            title_label = QtWidgets.QLabel(title)
            value_label = QtWidgets.QLabel("—")
            value_label.setTextInteractionFlags(QtCore.Qt.TextSelectableByMouse)
            metrics_layout.addWidget(title_label, idx // 2, (idx % 2) * 2)
            metrics_layout.addWidget(value_label, idx // 2, (idx % 2) * 2 + 1)
            self.metric_labels[key] = value_label
        detail_layout.addWidget(metrics_box)

        history_button = QtWidgets.QPushButton("История хоста")
        history_button.clicked.connect(self.show_history_dialog)
        detail_layout.addWidget(history_button)

        detail_layout.addWidget(self.section_label("График задержек"))
        self.chart = QChart()
        self.chart.setBackgroundVisible(False)
        self.chart.setLegendVisible(False)
        self.chart_view = QChartView(self.chart)
        self.chart_view.setRenderHint(QPainter.Antialiasing)
        detail_layout.addWidget(self.chart_view, 1)

    def build_stats_tab(self):
        layout = QtWidgets.QVBoxLayout(self.stats_tab)
        toolbar = QtWidgets.QHBoxLayout()
        layout.addLayout(toolbar)
        refresh_button = QtWidgets.QPushButton("Обновить статистику")
        refresh_button.clicked.connect(self.refresh_stats)
        toolbar.addWidget(refresh_button)
        toolbar.addStretch()

        self.stats_table = QtWidgets.QTableWidget(0, 7)
        self.stats_table.setHorizontalHeaderLabels(
            ["Мониторинг", "", "Группа", "Всего", "Доступно", "Недоступно", "Неизвестно"]
        )
        self.stats_table.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.stats_table.verticalHeader().setVisible(False)
        self.stats_table.horizontalHeader().setStretchLastSection(True)
        layout.addWidget(self.stats_table)

    def create_button(self, text, callback):
        button = QtWidgets.QPushButton(text)
        button.clicked.connect(callback)
        return button

    def section_label(self, text):
        label = QtWidgets.QLabel(text)
        font = label.font()
        font.setBold(True)
        label.setFont(font)
        return label

    def clear_search(self):
        self.search_input.clear()
        self.refresh_table()

    def load_groups(self):
        groups = self.db.get_groups()
        self.group_combo.clear()
        self.group_combo.addItems(groups)
        if groups:
            self.current_group = groups[0]
            self.group_combo.setCurrentText(self.current_group)
        else:
            self.current_group = None
            self.clear_table()
        self.refresh_stats()

    def on_group_select(self, group_name):
        self.current_group = group_name
        self.selected_host = None
        self.reset_detail_panel()
        self.refresh_table()

    def reset_detail_panel(self):
        for label in self.detail_labels.values():
            label.setText("—")
        for label in self.metric_labels.values():
            label.setText("—")
        self.clear_chart()

    def clear_table(self):
        self.host_table.setRowCount(0)

    def match_filter(self, status, offline_since):
        filter_value = self.filter_combo.currentText()
        if filter_value == "all":
            return True
        if filter_value == "unknown":
            return status == "Unknown"
        if filter_value == "online":
            return status == "Online"
        if filter_value.startswith("offline"):
            if status != "Offline":
                return False
            if not offline_since:
                return filter_value == "offline_lt_1h"
            try:
                offline_time = datetime.strptime(offline_since, "%Y-%m-%d %H:%M:%S")
                diff = (datetime.now() - offline_time).total_seconds()
            except ValueError:
                diff = 3600
            if filter_value == "offline_lt_1h":
                return diff < 3600
            if filter_value == "offline_ge_1h":
                return diff >= 3600
        return True

    def refresh_table(self):
        self.host_table.setRowCount(0)
        if not self.current_group:
            return

        search_text = self.search_input.text().strip().lower()
        hosts = self.db.get_hosts(self.current_group)
        self.host_row_map = {}

        for host in hosts:
            address, desc, subgroup, interval_sec, last_ping_time, offline_since = host
            if search_text:
                haystack = " ".join(
                    [str(address), str(desc or ""), str(subgroup or "")]
                ).lower()
                if search_text not in haystack:
                    continue

            status = self.db.get_last_status(self.current_group, address)
            if not self.match_filter(status, offline_since):
                continue

            self.host_status_cache[address] = (status, offline_since)

            interval_value = interval_sec if interval_sec else self.default_interval_seconds
            recent = self.db.get_recent_results(self.current_group, address, limit=1)
            last_latency = recent[-1][2] if recent else None
            latency_ms = f"{last_latency * 1000:.0f}" if last_latency else "-"
            last_check_display = last_ping_time or "-"

            row = self.host_table.rowCount()
            self.host_table.insertRow(row)
            self.host_table.setItem(row, 0, QtWidgets.QTableWidgetItem("●"))
            self.host_table.setItem(row, 1, QtWidgets.QTableWidgetItem(address))
            self.host_table.setItem(row, 2, QtWidgets.QTableWidgetItem(desc or ""))
            self.host_table.setItem(row, 3, QtWidgets.QTableWidgetItem(subgroup or ""))
            self.host_table.setItem(row, 4, QtWidgets.QTableWidgetItem(str(interval_value)))
            self.host_table.setItem(row, 5, QtWidgets.QTableWidgetItem(latency_ms))
            self.host_table.setItem(row, 6, QtWidgets.QTableWidgetItem(last_check_display))
            self.host_row_map[address] = row

    def refresh_stats(self):
        self.stats_table.setRowCount(0)
        active_map = self.db.get_groups_with_status()
        stats = self.db.get_all_hosts_stats()

        for entry in stats:
            row = self.stats_table.rowCount()
            self.stats_table.insertRow(row)

            checkbox = QtWidgets.QCheckBox()
            checkbox.setChecked(active_map.get(entry["group"], False))
            checkbox.stateChanged.connect(
                lambda state, group=entry["group"]: self.db.set_group_active(group, state == QtCore.Qt.Checked)
            )
            checkbox.stateChanged.connect(lambda _: self.refresh_stats())
            self.stats_table.setCellWidget(row, 0, checkbox)
            self.stats_table.setItem(row, 1, QtWidgets.QTableWidgetItem("●"))
            self.stats_table.setItem(row, 2, QtWidgets.QTableWidgetItem(entry["group"]))
            self.stats_table.setItem(row, 3, QtWidgets.QTableWidgetItem(str(entry["total"])))
            self.stats_table.setItem(row, 4, QtWidgets.QTableWidgetItem(str(entry["online"])))
            self.stats_table.setItem(row, 5, QtWidgets.QTableWidgetItem(str(entry["offline"])))
            self.stats_table.setItem(row, 6, QtWidgets.QTableWidgetItem(str(entry["unknown"])))

    def update_overall_status(self):
        overall = self.db.get_overall_status()
        color_map = {
            "healthy": "#2ecc71",
            "warning": "#f39c12",
            "critical": "#e74c3c",
        }
        self.overall_status_label.setStyleSheet(
            f"color: {color_map.get(overall, '#95a5a6')};"
        )

    def show_single_dialog(self, current_dialog, builder):
        if current_dialog is not None and current_dialog.isVisible():
            current_dialog.raise_()
            current_dialog.activateWindow()
            return current_dialog
        dialog = builder()
        dialog.show()
        return dialog

    def show_create_group_dialog(self):
        def build():
            dialog = QtWidgets.QDialog(self)
            dialog.setWindowTitle("Новая группа")
            dialog.setFixedSize(300, 160)
            layout = QtWidgets.QVBoxLayout(dialog)
            layout.addWidget(QtWidgets.QLabel("Имя группы:"))
            name_input = QtWidgets.QLineEdit()
            layout.addWidget(name_input)
            button = QtWidgets.QPushButton("Создать")
            layout.addWidget(button)

            def save():
                name = name_input.text().strip()
                if not name:
                    return
                try:
                    self.db.create_group(name)
                    self.load_groups()
                    self.group_combo.setCurrentText(name)
                    dialog.accept()
                except ValueError as exc:
                    QtWidgets.QMessageBox.warning(dialog, "Ошибка", str(exc))

            button.clicked.connect(save)
            return dialog

        self.create_group_dialog_ref = self.show_single_dialog(
            self.create_group_dialog_ref, build
        )

    def show_add_host_dialog(self):
        if not self.current_group:
            QtWidgets.QMessageBox.warning(self, "Внимание", "Выберите группу.")
            return

        def build():
            dialog = QtWidgets.QDialog(self)
            dialog.setWindowTitle("Добавить хост")
            dialog.setFixedSize(350, 340)
            layout = QtWidgets.QFormLayout(dialog)
            addr_input = QtWidgets.QLineEdit()
            desc_input = QtWidgets.QLineEdit()
            sub_input = QtWidgets.QLineEdit()
            interval_input = QtWidgets.QLineEdit(str(self.default_interval_seconds))
            layout.addRow("Адрес (IP/Домен):", addr_input)
            layout.addRow("Описание:", desc_input)
            layout.addRow("Подгруппа:", sub_input)
            layout.addRow("Интервал пинга (секунды):", interval_input)
            save_button = QtWidgets.QPushButton("Сохранить")
            layout.addRow(save_button)

            def save():
                try:
                    interval_seconds = int(interval_input.text())
                    if interval_seconds <= 0:
                        raise ValueError("Интервал должен быть больше 0 секунд.")
                    self.db.add_host(
                        self.current_group,
                        addr_input.text(),
                        desc_input.text(),
                        sub_input.text() or None,
                        interval_seconds,
                    )
                    self.refresh_table()
                    dialog.accept()
                except ValueError as exc:
                    QtWidgets.QMessageBox.warning(dialog, "Ошибка", str(exc))
                except sqlite3.Error as exc:
                    QtWidgets.QMessageBox.warning(dialog, "Ошибка БД", str(exc))

            save_button.clicked.connect(save)
            return dialog

        self.add_host_dialog = self.show_single_dialog(self.add_host_dialog, build)

    def show_edit_host_dialog(self):
        if not self.current_group:
            QtWidgets.QMessageBox.warning(self, "Внимание", "Выберите группу.")
            return

        selected_items = self.host_table.selectedItems()
        if not selected_items:
            QtWidgets.QMessageBox.warning(self, "Внимание", "Выберите хост.")
            return
        address = selected_items[1].text()
        hosts = self.db.get_hosts(self.current_group)
        host_map = {h[0]: h for h in hosts}
        host_data = host_map.get(address)
        if not host_data:
            return

        def build():
            dialog = QtWidgets.QDialog(self)
            dialog.setWindowTitle("Редактировать хост")
            dialog.setFixedSize(350, 340)
            layout = QtWidgets.QFormLayout(dialog)
            addr_input = QtWidgets.QLineEdit(host_data[0])
            addr_input.setEnabled(False)
            desc_input = QtWidgets.QLineEdit(host_data[1] or "")
            sub_input = QtWidgets.QLineEdit(host_data[2] or "")
            interval_seconds = host_data[3] if host_data[3] else self.default_interval_seconds
            interval_input = QtWidgets.QLineEdit(str(interval_seconds))
            layout.addRow("Адрес (IP/Домен):", addr_input)
            layout.addRow("Описание:", desc_input)
            layout.addRow("Подгруппа:", sub_input)
            layout.addRow("Интервал пинга (секунды):", interval_input)
            save_button = QtWidgets.QPushButton("Сохранить")
            layout.addRow(save_button)

            def save():
                try:
                    interval_value = int(interval_input.text())
                    if interval_value <= 0:
                        raise ValueError("Интервал должен быть больше 0 секунд.")
                    self.db.update_host(
                        self.current_group,
                        host_data[0],
                        desc_input.text(),
                        sub_input.text() or None,
                        interval_value,
                    )
                    self.refresh_table()
                    dialog.accept()
                except ValueError as exc:
                    QtWidgets.QMessageBox.warning(dialog, "Ошибка", str(exc))
                except sqlite3.Error as exc:
                    QtWidgets.QMessageBox.warning(dialog, "Ошибка БД", str(exc))

            save_button.clicked.connect(save)
            return dialog

        self.edit_host_dialog = self.show_single_dialog(self.edit_host_dialog, build)

    def show_settings_dialog(self):
        def build():
            dialog = QtWidgets.QDialog(self)
            dialog.setWindowTitle("Настройки")
            dialog.setFixedSize(360, 320)
            layout = QtWidgets.QFormLayout(dialog)
            default_interval_input = QtWidgets.QLineEdit(str(self.default_interval_seconds))
            monitor_interval_input = QtWidgets.QLineEdit(str(self.monitor_interval))
            ping_timeout_input = QtWidgets.QLineEdit(str(self.ping_timeout))
            batch_size_input = QtWidgets.QLineEdit(str(self.batch_size))
            max_workers_input = QtWidgets.QLineEdit(str(self.max_workers))
            layout.addRow("Интервал по умолчанию (сек)", default_interval_input)
            layout.addRow("Интервал цикла мониторинга (сек)", monitor_interval_input)
            layout.addRow("Таймаут пинга (сек)", ping_timeout_input)
            layout.addRow("Размер пачки пингов", batch_size_input)
            layout.addRow("Макс. потоков пинга", max_workers_input)
            save_button = QtWidgets.QPushButton("Сохранить")
            layout.addRow(save_button)

            def save():
                try:
                    default_interval = int(default_interval_input.text())
                    monitor_interval = float(monitor_interval_input.text())
                    ping_timeout = float(ping_timeout_input.text())
                    batch_size = int(batch_size_input.text())
                    max_workers = int(max_workers_input.text())

                    if default_interval <= 0 or monitor_interval <= 0 or ping_timeout <= 0:
                        raise ValueError("Интервалы и таймаут должны быть больше 0.")
                    if batch_size <= 0 or max_workers <= 0:
                        raise ValueError("Размеры должны быть больше 0.")

                    self.db.set_setting("default_interval_seconds", default_interval)
                    self.db.set_setting("monitor_interval_seconds", monitor_interval)
                    self.db.set_setting("ping_timeout_seconds", ping_timeout)
                    self.db.set_setting("ping_batch_size", batch_size)
                    self.db.set_setting("max_ping_workers", max_workers)

                    self.load_settings()
                    dialog.accept()
                except ValueError as exc:
                    QtWidgets.QMessageBox.warning(dialog, "Ошибка", str(exc))

            save_button.clicked.connect(save)
            return dialog

        self.settings_dialog = self.show_single_dialog(self.settings_dialog, build)

    def show_history_dialog(self):
        if not self.selected_host:
            QtWidgets.QMessageBox.warning(self, "Внимание", "Выберите хост.")
            return

        group, address = self.selected_host

        def build():
            dialog = QtWidgets.QDialog(self)
            dialog.setWindowTitle(f"История хоста {address}")
            dialog.resize(700, 500)
            layout = QtWidgets.QVBoxLayout(dialog)
            layout.addWidget(QtWidgets.QLabel(f"Группа: {group}"))

            change_box = QtWidgets.QGroupBox("Смена статуса")
            change_layout = QtWidgets.QVBoxLayout(change_box)
            change_table = QtWidgets.QTableWidget(0, 2)
            change_table.setHorizontalHeaderLabels(["Время", "Статус"])
            change_table.verticalHeader().setVisible(False)
            change_table.horizontalHeader().setStretchLastSection(True)
            change_layout.addWidget(change_table)
            layout.addWidget(change_box)

            history_box = QtWidgets.QGroupBox("Полная история")
            history_layout = QtWidgets.QVBoxLayout(history_box)
            history_table = QtWidgets.QTableWidget(0, 3)
            history_table.setHorizontalHeaderLabels(["Время", "Статус", "Задержка (мс)"])
            history_table.verticalHeader().setVisible(False)
            history_table.horizontalHeader().setStretchLastSection(True)
            history_layout.addWidget(history_table)
            layout.addWidget(history_box)

            results = self.db.get_recent_results(group, address, limit=200)
            last_status = None
            for timestamp, status, latency in results:
                latency_ms = f"{latency * 1000:.0f}" if latency else "-"
                row = history_table.rowCount()
                history_table.insertRow(row)
                history_table.setItem(row, 0, QtWidgets.QTableWidgetItem(timestamp))
                history_table.setItem(row, 1, QtWidgets.QTableWidgetItem(status))
                history_table.setItem(row, 2, QtWidgets.QTableWidgetItem(latency_ms))

                if status != last_status:
                    change_row = change_table.rowCount()
                    change_table.insertRow(change_row)
                    change_table.setItem(change_row, 0, QtWidgets.QTableWidgetItem(timestamp))
                    change_table.setItem(change_row, 1, QtWidgets.QTableWidgetItem(status))
                    last_status = status

            return dialog

        self.history_dialog = self.show_single_dialog(self.history_dialog, build)

    def delete_group(self):
        if not self.current_group:
            return
        reply = QtWidgets.QMessageBox.question(
            self,
            "Подтверждение",
            f"Удалить группу '{self.current_group}'?",
        )
        if reply == QtWidgets.QMessageBox.Yes:
            self.db.delete_group(self.current_group)
            self.current_group = None
            self.load_groups()

    def remove_host(self):
        selected_items = self.host_table.selectedItems()
        if not selected_items:
            return
        address = selected_items[1].text()
        self.db.remove_host(self.current_group, address)
        self.refresh_table()

    def export_data(self):
        if not self.current_group:
            return
        if not HAS_OPENPYXL:
            QtWidgets.QMessageBox.warning(self, "Ошибка", "openpyxl не установлен.")
            return
        filename, _ = QtWidgets.QFileDialog.getSaveFileName(
            self, "Экспорт", "", "Excel Files (*.xlsx)"
        )
        if filename:
            wb = Workbook()
            ws = wb.active
            ws.title = "hosts"
            ws.append(["address", "description", "subgroup", "ping_interval_sec"])
            hosts = self.db.get_hosts(self.current_group)
            for h in hosts:
                interval_sec = h[3] if h[3] else self.default_interval_seconds
                ws.append([h[0], h[1], h[2] or "", interval_sec])
            for col in ws.columns:
                max_length = 0
                for cell in col:
                    cell_value = str(cell.value) if cell.value is not None else ""
                    max_length = max(max_length, len(cell_value))
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[get_column_letter(col[0].column)].width = adjusted_width
            wb.save(filename)

    def import_data(self):
        if not self.current_group:
            QtWidgets.QMessageBox.warning(self, "Внимание", "Выберите группу.")
            return
        if not HAS_OPENPYXL:
            QtWidgets.QMessageBox.warning(self, "Ошибка", "openpyxl не установлен.")
            return
        filename, _ = QtWidgets.QFileDialog.getOpenFileName(
            self, "Импорт", "", "Excel Files (*.xlsx)"
        )
        if filename:
            wb = load_workbook(filename)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return
            headers = [str(cell).strip().lower() if cell else "" for cell in rows[0]]
            interval_is_minutes = "ping_interval_min" in headers
            for row in rows[1:]:
                if row and len(row) >= 2:
                    addr = row[0]
                    desc = row[1]
                    sub = row[2] if len(row) > 2 else None
                    interval_value = (
                        int(row[3]) if len(row) > 3 and row[3] else self.default_interval_seconds
                    )
                    interval_sec = interval_value * 60 if interval_is_minutes else interval_value
                    if addr:
                        self.db.add_host(self.current_group, addr, desc, sub, interval_sec)
            self.refresh_table()

    def on_host_select(self):
        selected_items = self.host_table.selectedItems()
        if not selected_items:
            return
        address = selected_items[1].text()
        self.selected_host = (self.current_group, address)
        self.update_detail_panel(address)
        self.update_chart()

    def update_detail_panel(self, address):
        if not self.current_group or not address:
            return
        hosts = self.db.get_hosts(self.current_group)
        host_map = {h[0]: h for h in hosts}
        host_data = host_map.get(address)
        if not host_data:
            return

        address, desc, subgroup, interval_sec, last_ping_time, _ = host_data
        status = self.db.get_last_status(self.current_group, address)
        recent_results = self.db.get_recent_results(self.current_group, address, limit=50)
        latest_latency = recent_results[-1][2] if recent_results else None
        last_status, last_change_time = self.db.get_last_status_change(
            self.current_group, address
        )

        latencies = [row[2] for row in recent_results if row[2] is not None]
        loss_count = sum(1 for row in recent_results if row[1] == "Offline")
        total_count = len(recent_results)
        jitter = None
        if len(latencies) > 1:
            diffs = [abs(latencies[i] - latencies[i - 1]) for i in range(1, len(latencies))]
            jitter = sum(diffs) / len(diffs)

        self.detail_labels["Адрес"].setText(address)
        self.detail_labels["Описание"].setText(desc or "—")
        self.detail_labels["Подгруппа"].setText(subgroup or "—")
        interval_value = interval_sec if interval_sec else self.default_interval_seconds
        self.detail_labels["Интервал"].setText(f"{interval_value} сек")
        self.detail_labels["Статус"].setText(status)
        self.detail_labels["Задержка"].setText(
            f"{latest_latency * 1000:.0f} мс" if latest_latency else "—"
        )
        self.detail_labels["Последняя проверка"].setText(last_ping_time or "—")
        if last_status and last_change_time:
            self.detail_labels["Последняя смена"].setText(
                f"{last_status} в {last_change_time}"
            )
        else:
            self.detail_labels["Последняя смена"].setText("—")

        if latencies:
            self.metric_labels["min"].setText(f"{min(latencies) * 1000:.0f} мс")
            self.metric_labels["max"].setText(f"{max(latencies) * 1000:.0f} мс")
            self.metric_labels["avg"].setText(
                f"{(sum(latencies) / len(latencies)) * 1000:.0f} мс"
            )
        else:
            self.metric_labels["min"].setText("—")
            self.metric_labels["max"].setText("—")
            self.metric_labels["avg"].setText("—")

        self.metric_labels["jitter"].setText(
            f"{jitter * 1000:.0f} мс" if jitter is not None else "—"
        )
        if total_count > 0:
            self.metric_labels["loss"].setText(f"{(loss_count / total_count) * 100:.1f}%")
        else:
            self.metric_labels["loss"].setText("—")

    def clear_chart(self):
        self.chart.removeAllSeries()
        self.chart.createDefaultAxes()

    def update_chart(self):
        self.chart.removeAllSeries()
        if not self.selected_host:
            return
        group, address = self.selected_host
        data = self.db.get_recent_results(group, address, limit=50)
        if not data:
            return

        line_series = QLineSeries()
        point_series = QScatterSeries()
        point_series.setMarkerSize(6.0)
        offline_series = QScatterSeries()
        offline_series.setMarkerSize(8.0)
        offline_series.setColor(QtCore.Qt.red)

        x_values = list(range(len(data)))
        latencies_ms = [row[2] * 1000 for row in data if row[2] is not None]
        min_latency = min(latencies_ms) if latencies_ms else 0
        max_latency = max(latencies_ms) if latencies_ms else 1

        for idx, (_, status, latency) in enumerate(data):
            if latency is None:
                offline_series.append(idx, min_latency)
                continue
            latency_ms = latency * 1000
            line_series.append(idx, latency_ms)
            point_series.append(idx, latency_ms)

        self.chart.addSeries(line_series)
        self.chart.addSeries(point_series)
        self.chart.addSeries(offline_series)

        axis_x = QValueAxis()
        axis_x.setRange(0, max(len(data) - 1, 1))
        axis_x.setTickCount(5)
        axis_x.setLabelFormat("%d")
        axis_x.setTitleText("Точки")

        axis_y = QValueAxis()
        axis_y.setRange(min_latency, max_latency if max_latency > min_latency else min_latency + 1)
        axis_y.setLabelFormat("%d")
        axis_y.setTitleText("мс")

        self.chart.addAxis(axis_x, QtCore.Qt.AlignBottom)
        self.chart.addAxis(axis_y, QtCore.Qt.AlignLeft)
        line_series.attachAxis(axis_x)
        line_series.attachAxis(axis_y)
        point_series.attachAxis(axis_x)
        point_series.attachAxis(axis_y)
        offline_series.attachAxis(axis_x)
        offline_series.attachAxis(axis_y)

    def toggle_monitoring(self):
        if self.monitoring:
            self.monitoring = False
            self.btn_start.setText("Старт")
            self.status_label.setText("Мониторинг остановлен")
        else:
            if not self.db.get_active_groups():
                QtWidgets.QMessageBox.warning(
                    self,
                    "Внимание",
                    "Нет активных групп для мониторинга. Включите группы на вкладке 'Статистика'.",
                )
                return
            self.monitoring = True
            self.btn_start.setText("Стоп")
            self.status_label.setText("Мониторинг запущен...")
            self.monitor_thread = threading.Thread(target=self.monitoring_loop, daemon=True)
            self.monitor_thread.start()

    def is_due(self, last_ping_time, interval_seconds, now):
        if not last_ping_time:
            return True
        try:
            last_time = datetime.strptime(last_ping_time, "%Y-%m-%d %H:%M:%S")
        except ValueError:
            return True
        return (now - last_time).total_seconds() >= interval_seconds

    def chunk_hosts(self, hosts, chunk_size):
        for i in range(0, len(hosts), chunk_size):
            yield hosts[i : i + chunk_size]

    def ping_hosts_batch(self, hosts_batch):
        entries = []
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with concurrent.futures.ThreadPoolExecutor(max_workers=self.max_workers) as executor:
            future_to_host = {
                executor.submit(ping_host_subprocess, host[1], self.ping_timeout): host
                for host in hosts_batch
            }
            for future in concurrent.futures.as_completed(future_to_host):
                host = future_to_host[future]
                group_name, address, subgroup = host
                try:
                    success, latency = future.result()
                    status = "Online" if success else "Offline"
                    entries.append(
                        {
                            "timestamp": timestamp,
                            "group_name": group_name,
                            "subgroup": subgroup,
                            "address": address,
                            "status": status,
                            "latency": latency,
                        }
                    )
                except Exception:
                    entries.append(
                        {
                            "timestamp": timestamp,
                            "group_name": group_name,
                            "subgroup": subgroup,
                            "address": address,
                            "status": "Offline",
                            "latency": None,
                        }
                    )
        return entries

    def monitoring_loop(self):
        while self.monitoring:
            groups = self.db.get_active_groups()
            if not groups:
                time.sleep(self.monitor_interval)
                continue

            now = datetime.now()
            due_hosts = []
            for group in groups:
                hosts = self.db.get_hosts(group)
                for host in hosts:
                    address, _, subgroup, interval_sec, last_ping_time, _ = host
                    interval_sec = interval_sec or self.default_interval_seconds
                    if self.is_due(last_ping_time, interval_sec, now):
                        due_hosts.append((group, address, subgroup))

            if due_hosts:
                for batch in self.chunk_hosts(due_hosts, self.batch_size):
                    if not self.monitoring:
                        break
                    results = self.ping_hosts_batch(batch)
                    self.db.log_results_batch(results)
                    for result in results:
                        self.ui_queue.put(
                            {
                                "type": "row",
                                "group": result["group_name"],
                                "address": result["address"],
                                "status": result["status"],
                                "latency": result["latency"],
                                "timestamp": result["timestamp"],
                            }
                        )

            now_ts = time.time()
            if now_ts - self.last_stats_refresh >= STATS_REFRESH_INTERVAL:
                self.ui_queue.put({"type": "stats"})
                self.ui_queue.put({"type": "overall"})
                self.last_stats_refresh = now_ts

            time.sleep(self.monitor_interval)

    def update_row(self, group, address, status, latency, timestamp):
        if self.current_group != group:
            return
        row = self.host_row_map.get(address)
        if row is None:
            return

        offline_since = self.host_status_cache.get(address, ("Unknown", None))[1]
        if status == "Offline" and offline_since is None:
            offline_since = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        elif status == "Online":
            offline_since = None

        self.host_status_cache[address] = (status, offline_since)
        latency_str = f"{latency * 1000:.0f}" if latency else "-"

        self.host_table.setItem(row, 5, QtWidgets.QTableWidgetItem(latency_str))
        self.host_table.setItem(row, 6, QtWidgets.QTableWidgetItem(timestamp))

        if self.selected_host == (group, address):
            self.update_detail_panel(address)
            self.update_chart()

    def process_ui_queue(self):
        while True:
            try:
                event = self.ui_queue.get_nowait()
            except queue.Empty:
                break
            if event["type"] == "row":
                self.update_row(
                    event["group"],
                    event["address"],
                    event["status"],
                    event["latency"],
                    event["timestamp"],
                )
            elif event["type"] == "stats":
                self.refresh_stats()
            elif event["type"] == "overall":
                self.update_overall_status()


if __name__ == "__main__":
    app = QtWidgets.QApplication([])
    window = PingApp()
    window.show()
    app.exec()
